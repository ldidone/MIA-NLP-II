"""Document loaders for CV files (PDF, DOCX, TXT)."""

from __future__ import annotations

import io
from pathlib import Path

from pypdf import PdfReader
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class UnsupportedFileTypeError(ValueError):
    """Raised when the uploaded file extension is not supported."""


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    paragraphs = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from a CV file.

    Args:
        file_bytes: Raw bytes of the uploaded file.
        filename: Original filename, used to detect extension.

    Returns:
        A unicode string containing the extracted text.

    Raises:
        UnsupportedFileTypeError: When the extension is not in SUPPORTED_EXTENSIONS.
    """
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = _extract_pdf(file_bytes)
    elif ext == ".docx":
        text = _extract_docx(file_bytes)
    elif ext == ".txt":
        text = _extract_txt(file_bytes)
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type: {ext!r}. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )
    return _normalize_whitespace(text)


def _normalize_whitespace(text: str) -> str:
    lines = [line.strip() for line in text.splitlines()]
    cleaned = "\n".join(line for line in lines if line)
    return cleaned.strip()
